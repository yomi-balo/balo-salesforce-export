import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# --- Style constants ---
HEADER_FILL = PatternFill("solid", fgColor="2F5496")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

SEPARATOR_FILL = PatternFill("solid", fgColor="D6E4F0")
SEPARATOR_FONT = Font(bold=True, italic=True, size=11, color="2F5496")

HELPER_COL_FILL = PatternFill("solid", fgColor="E8E8E8")
HELPER_HEADER_FILL = PatternFill("solid", fgColor="808080")

FLAG_FILL = PatternFill("solid", fgColor="FFD966")
FLAG_FONT = Font(bold=True, color="CC0000")

THIN_BORDER = Border(
    bottom=Side(style="thin", color="D0D0D0"),
)

TAB_COLORS = {
    "Accounts": "4472C4",
    "Prospects & Contacts": "ED7D31",
    "Opportunities - Cases": "A9D18E",
    "Opportunities - Projects": "FFD966",
    "Project Experts": "9DC3E6",
    "Consultations": "F4B183",
}

HELPER_COLUMNS = {"_group", "_contact_type", "_meeting_type", "_source_stage"}


def write_workbook(sheets_data, output_path, warnings=None):
    """Generate the full XLSX workbook.

    Args:
        sheets_data: dict of {sheet_name: (columns, rows)} from transformer
        output_path: file path for the .xlsx
        warnings: optional list of warning strings
    """
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    wb = Workbook()

    # Sheet 1: README
    _write_readme(wb.active, warnings or [])

    # Data sheets
    sheet_order = [
        "Accounts",
        "Prospects & Contacts",
        "Opportunities - Cases",
        "Opportunities - Projects",
        "Project Experts",
        "Consultations",
    ]

    for sheet_name in sheet_order:
        if sheet_name not in sheets_data:
            continue
        columns, rows = sheets_data[sheet_name]
        ws = wb.create_sheet(title=sheet_name)
        ws.sheet_properties.tabColor = TAB_COLORS.get(sheet_name, "000000")
        _write_data_sheet(ws, columns, rows)

    try:
        wb.save(output_path)
    except PermissionError:
        raise PermissionError(
            f"Cannot write to {output_path}. Is the file open in Excel? Close it and try again."
        )


def _write_readme(ws, warnings):
    ws.title = "README"
    ws.sheet_properties.tabColor = "333333"
    ws.column_dimensions["A"].width = 100

    row = 1
    mono = Font(name="Courier New", size=10)
    title_font = Font(bold=True, size=16, color="2F5496")
    section_font = Font(bold=True, size=13, color="2F5496")
    body_font = Font(size=11)

    ws.cell(row=row, column=1, value="Balo → Salesforce Test Data Export").font = title_font
    row += 1
    ws.cell(row=row, column=1, value=f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").font = Font(size=10, color="666666")
    row += 2

    # How to use
    ws.cell(row=row, column=1, value="HOW TO USE THIS TEST DATA").font = section_font
    row += 1
    for line in [
        "Each sheet maps to a Salesforce API endpoint. Column names = SF payload field names.",
        "Use the _group column to find related records across sheets.",
        "Columns prefixed with _ are for reference only — do not send to SF.",
        "",
    ]:
        ws.cell(row=row, column=1, value=line).font = body_font
        row += 1

    # Endpoint reference
    ws.cell(row=row, column=1, value="ENDPOINT REFERENCE").font = section_font
    row += 1
    endpoints = [
        "Accounts        → PATCH /sobjects/Account/Balo_Id__c/{id}",
        "Prospects        → POST  /services/apexrest/Prospect/",
        "Contact (Client) → PATCH /sobjects/Contact/Balo_Id__c/{baloId}",
        "Contact (Expert) → PATCH /sobjects/Contact/Balo_Id__c/{baloId}",
        "Opp (Case)       → PATCH /sobjects/Opportunity/Balo_Id__c/{id}",
        "Opp (Project)    → PATCH /sobjects/Opportunity/Balo_Id__c/{id}",
        "Project Expert   → PATCH /sobjects/Project__c/Balo_Id__c/{compositeId}",
        "Consultation     → PATCH /sobjects/Consultation__c/Balo_Id__c/{meetingId}",
    ]
    for ep in endpoints:
        ws.cell(row=row, column=1, value=ep).font = mono
        row += 1
    row += 1

    # Cross-reference
    ws.cell(row=row, column=1, value="CROSS-REFERENCE GUIDE").font = section_font
    row += 1
    refs = [
        "Account (Accounts)    ←  baloAccountId in Prospects & Contacts",
        "Account (Accounts)    ←  Account.Balo_Id__c in all sheets",
        "Contact (Prospects)   ←  Primary_Contact__r.Balo_Id__c in Opportunities",
        "Contact (Prospects)   ←  Expert__r.Balo_Id__c in Cases & Consultations",
        "Opp Case (Cases)      ←  Opportunity__r.Balo_Id__c in Consultations",
        "Opp Project (Projects)←  Opportunity__r.Balo_Id__c in Project Experts",
        "Project Expert        ←  Project__r.Balo_Id__c in Consultations (project meetings)",
    ]
    for ref in refs:
        ws.cell(row=row, column=1, value=ref).font = mono
        row += 1
    row += 1

    # Example flow
    ws.cell(row=row, column=1, value="EXAMPLE FLOW — Test a full client journey:").font = section_font
    row += 1
    steps = [
        "1. POST Prospect: take a client row from 'Prospects & Contacts'",
        "2. PATCH Account: take the matching row from 'Accounts'",
        "3. PATCH Contact: take the same client row's Contact columns",
        "4. PATCH Opportunity (Case): take a row from 'Opportunities - Cases'",
        "5. PATCH Consultation: take a consultation row with matching Opportunity__r",
    ]
    for step in steps:
        ws.cell(row=row, column=1, value=step).font = body_font
        row += 1
    row += 1

    # Flags for Nick
    ws.cell(row=row, column=1, value="FLAGS FOR NICK").font = section_font
    row += 1
    flags = [
        "RecordTypeId for Project Opportunity: confirm 012On00000L9nILIAZ vs 012On00000L9nDVIAZ",
        "Does Project_Tag__c accept multiple values? If so, what separator?",
        "Consultation endpoint: Nick flagged as still needing review (March 23, 2026)",
    ]
    for flag in flags:
        cell = ws.cell(row=row, column=1, value=flag)
        cell.font = FLAG_FONT
        cell.fill = FLAG_FILL
        row += 1
    row += 1

    # Warnings from export
    if warnings:
        ws.cell(row=row, column=1, value="EXPORT WARNINGS").font = section_font
        row += 1
        for w in warnings[:50]:  # cap at 50
            ws.cell(row=row, column=1, value=w).font = Font(size=10, color="CC6600")
            row += 1


def _write_data_sheet(ws, columns, rows):
    helper_col_indices = set()

    # --- Write header row ---
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        if col_name in HELPER_COLUMNS:
            cell.fill = HELPER_HEADER_FILL
            helper_col_indices.add(col_idx)
        else:
            cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT

    # Freeze header
    ws.freeze_panes = "A2"

    # --- Write data rows ---
    data_row = 2
    for row_data in rows:
        if row_data.get("_is_separator"):
            # Separator row
            label = row_data.get("_label", "")
            cell = ws.cell(row=data_row, column=1, value=label)
            cell.font = SEPARATOR_FONT
            cell.fill = SEPARATOR_FILL
            # Fill entire row with separator color
            for col_idx in range(1, len(columns) + 1):
                ws.cell(row=data_row, column=col_idx).fill = SEPARATOR_FILL
            # Merge across all columns
            if len(columns) > 1:
                ws.merge_cells(
                    start_row=data_row, start_column=1,
                    end_row=data_row, end_column=len(columns)
                )
            data_row += 1
            continue

        # Normal data row
        for col_idx, col_name in enumerate(columns, 1):
            value = row_data.get(col_name, "")

            # Format booleans
            if isinstance(value, bool):
                value = "TRUE" if value else "FALSE"

            cell = ws.cell(row=data_row, column=col_idx, value=value)

            # Helper column styling
            if col_idx in helper_col_indices:
                cell.fill = HELPER_COL_FILL

            # FLAG cell highlighting
            if isinstance(value, str) and value.startswith("FLAG:"):
                cell.fill = FLAG_FILL
                cell.font = Font(bold=True, size=10)

            # Light bottom border for readability
            cell.border = THIN_BORDER

            # Right-align numbers
            if isinstance(value, (int, float)):
                cell.alignment = Alignment(horizontal="right")

        data_row += 1

    # --- Auto-filter ---
    if data_row > 2:
        last_col_letter = get_column_letter(len(columns))
        ws.auto_filter.ref = f"A1:{last_col_letter}{data_row - 1}"

    # --- Auto-width columns ---
    for col_idx in range(1, len(columns) + 1):
        max_len = len(str(columns[col_idx - 1]))
        for row_idx in range(2, min(data_row, 50)):  # sample first 50 rows
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            if cell_val:
                max_len = max(max_len, len(str(cell_val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 55)
